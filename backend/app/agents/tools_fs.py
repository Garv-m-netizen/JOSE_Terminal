"""Filesystem and document tools (approved folders only) with deterministic verifiers.

Registered into agents/custom_agent.py (IMPLS / VERIFIERS). All paths go through tools/fs_sandbox.resolve.
"""
from __future__ import annotations

import asyncio
import re
import shutil
from pathlib import Path
from typing import Any

from ..services import artifact_service
from ..tools import fs_sandbox as sbx
from ..tools import template_fill as tfill
from ..tools.extraction import norm
from ..tools.file_tools import extract_text, sha256_file
from ..workflows.custom import resolve_args
from .base import AgentContext, TaskError
from .verification_agent import Check

BLOCKED_EXT = {".exe", ".bat", ".cmd", ".ps1", ".sh", ".dll", ".msi", ".scr", ".vbs", ".jar", ".com", ".lnk"}
MAX_IMPORT_BYTES = 20 * 1024 * 1024
KINDS = {".pdf": "pdf", ".docx": "docx", ".xlsx": "xlsx", ".csv": "csv", ".pptx": "pptx"}


def _ck(name: str, ok: bool, reason: str) -> Check:
    return Check(name, bool(ok), reason)


def _sandbox(fn):
    """Turn sandbox violations into non-retryable validation errors (a retry cannot make a path legal)."""
    def wrap(*a, **k):
        try:
            return fn(*a, **k)
        except sbx.SandboxError as e:
            raise TaskError(str(e), "validation", retryable=False) from e
        except OSError as e:
            raise TaskError(f"file system error: {e}", "file", retryable=False) from e
    return wrap


def _text_of(v: Any) -> str:
    if isinstance(v, list):
        return "\n\n".join(map(str, v))
    if isinstance(v, dict):
        return str(v.get("text", v))
    return str(v)


# ------------------------------------------------------------------ tools

async def t_fs_list(ctx: AgentContext, a) -> dict:
    return await ctx.tool("fs.list", {"path": a.path}, lambda: asyncio.to_thread(_sandbox(sbx.list_dir), a.path),
                          lambda r: f"{r['count']} entries", simulated=False)


async def t_fs_search(ctx: AgentContext, a) -> dict:
    return await ctx.tool("fs.search", {"query": a.query, "path": a.path},
                          lambda: asyncio.to_thread(_sandbox(sbx.search), a.query, a.path, a.max_results, a.extensions),
                          lambda r: f"{r['count']} matches", simulated=False)


async def t_fs_read(ctx: AgentContext, a) -> dict:
    return await ctx.tool("fs.read", {"path": a.path}, lambda: asyncio.to_thread(_sandbox(sbx.read_text), a.path, a.max_chars),
                          lambda r: f"{r['chars']} characters read", simulated=False)


async def t_fs_import(ctx: AgentContext, a) -> dict:
    src = _sandbox(sbx.resolve)(a.path)
    if not src.is_file():
        raise TaskError(f"'{a.path}' is not a file", "validation", retryable=False)
    if src.suffix.lower() in BLOCKED_EXT:
        raise TaskError(f"files of type {src.suffix} are never imported", "validation", retryable=False)
    if src.stat().st_size > MAX_IMPORT_BYTES:
        raise TaskError("file is larger than 20 MB", "validation", retryable=False)
    from .custom_agent import safe_name

    name = safe_name(src.name)
    dest = artifact_service.run_dir(ctx.run_id) / name

    def copy() -> str:
        if dest.exists():
            dest.unlink()
        shutil.copy2(src, dest)
        return sha256_file(dest)

    async with ctx.activity(f"importing {src.name}"):
        digest = await ctx.tool("fs.import", {"path": a.path}, lambda: asyncio.to_thread(copy), lambda d: f"copied ({d[:12]})", simulated=False)
    art = artifact_service.register(ctx.run_id, ctx.task.id, dest, KINDS.get(dest.suffix.lower(), "other"), ctx.simulated)
    return {"artifact_id": art["id"], "filename": name, "sha256": digest, "source": sbx.rel(src), "size": dest.stat().st_size}


async def t_fs_write_text(ctx: AgentContext, a) -> dict:
    from .custom_agent import safe_name

    name = safe_name(a.filename)
    if not name.lower().endswith((".txt", ".md", ".csv", ".json")):
        raise TaskError("filename must end with .txt, .md, .csv or .json", "validation", retryable=False)
    text = _text_of(resolve_args(a.content, ctx.outputs))
    if not text.strip():
        raise TaskError("there is no content to write", "validation")
    path = artifact_service.run_dir(ctx.run_id) / name

    def write() -> None:
        path.write_text(text, encoding="utf-8", newline="\n")

    async with ctx.activity(f"writing {name}"):
        await ctx.tool("fs.write_text", {"file": name, "chars": len(text)}, lambda: asyncio.to_thread(write), lambda _: f"{name} written", simulated=False)
    art = artifact_service.register(ctx.run_id, ctx.task.id, path, KINDS.get(path.suffix.lower(), "other"), ctx.simulated)
    return {"artifact_id": art["id"], "filename": name, "chars": len(text), "expected_sha256": _sha_text(text)}


def _sha_text(text: str) -> str:
    import hashlib

    return hashlib.sha256(text.encode("utf-8")).hexdigest()


def _dest(destination: str, src: Path) -> Path:
    dst = _sandbox(sbx.resolve)(destination, must_exist=False)
    if dst.exists():
        raise TaskError(f"'{destination}' already exists; JOSE never overwrites files", "validation", retryable=False)
    if dst == src:
        raise TaskError("source and destination are the same", "validation", retryable=False)
    return dst


async def t_fs_copy(ctx: AgentContext, a) -> dict:
    src = _sandbox(sbx.resolve)(a.source)
    if not src.is_file():
        raise TaskError(f"'{a.source}' is not a file", "validation", retryable=False)
    dst = _dest(a.destination, src)
    before = sha256_file(src)

    def do() -> None:
        dst.parent.mkdir(parents=True, exist_ok=True)
        shutil.copy2(src, dst)

    async with ctx.activity(f"copying {src.name}"):
        await ctx.tool("fs.copy", {"source": a.source, "destination": a.destination}, lambda: asyncio.to_thread(do), lambda _: "copied", simulated=False)
    return {"source": sbx.rel(src), "destination": sbx.rel(dst), "sha256": before, "size": dst.stat().st_size}


async def t_fs_move(ctx: AgentContext, a) -> dict:
    src = _sandbox(sbx.resolve)(a.source)
    if not src.is_file():
        raise TaskError(f"'{a.source}' is not a file (it may already have been moved; JOSE will not guess)", "validation", retryable=False)
    dst = _dest(a.destination, src)
    before = sha256_file(src)
    size = src.stat().st_size

    def do() -> None:
        dst.parent.mkdir(parents=True, exist_ok=True)
        shutil.move(str(src), str(dst))

    async with ctx.activity(f"moving {src.name}"):
        await ctx.tool("fs.move", {"source": a.source, "destination": a.destination}, lambda: asyncio.to_thread(do), lambda _: "moved", simulated=False)
    return {"source": a.source, "destination": sbx.rel(dst), "sha256": before, "size": size}


async def t_doc_read(ctx: AgentContext, a) -> dict:
    from .custom_agent import find_upload

    path, shown = find_upload(ctx, a.file)

    def read() -> dict:
        ext = path.suffix.lower()
        if ext == ".xlsx":
            return {"kind": "xlsx", "text": tfill.read_xlsx_text(path)}
        try:
            r = extract_text(path)
        except Exception as e:  # noqa: BLE001
            raise TaskError(f"cannot read '{shown}': {e}", "file", retryable=False) from e
        if r.get("needs_ocr"):
            raise TaskError(f"'{shown}' is a scanned PDF without a text layer; OCR is not implemented", "file", retryable=False)
        return {"kind": r["kind"], "text": r["text"]}

    r = await ctx.tool("doc.read", {"file": shown}, lambda: asyncio.to_thread(read), lambda x: f"{len(x['text'])} characters", simulated=False)
    return {"file": shown, "kind": r["kind"], "chars": len(r["text"]), "text": r["text"][: a.max_chars], "truncated": len(r["text"]) > a.max_chars, "untrusted_content": True}


async def t_docx_compose(ctx: AgentContext, a) -> dict:
    import docx

    from .custom_agent import _rows, fmt_value, safe_name

    text = _text_of(resolve_args(a.body, ctx.outputs))
    name = safe_name(a.filename, ".docx")
    if not name.lower().endswith(".docx"):
        raise TaskError("filename must end with .docx", "validation", retryable=False)
    rows = _rows(resolve_args(a.table_source, ctx.outputs)) if a.table_source else []
    path = artifact_service.run_dir(ctx.run_id) / name

    def make() -> None:
        d = docx.Document()
        d.add_heading(a.title, 0)
        bullets: list[str] = []
        for para in re.split(r"\n\s*\n", text.strip()):
            for line in para.splitlines():
                if line.startswith("- "):
                    d.add_paragraph(line[2:], style="List Bullet")
                elif line.startswith("# "):
                    d.add_heading(line[2:], 1)
                elif line.strip():
                    d.add_paragraph(line)
        if rows:
            cols = list(rows[0])
            t = d.add_table(rows=1, cols=len(cols))
            t.style = "Table Grid"
            for i, c in enumerate(cols):
                t.rows[0].cells[i].text = str(c)
            for r in rows[:300]:
                cells = t.add_row().cells
                for i, c in enumerate(cols):
                    cells[i].text = fmt_value(r.get(c))
        d.save(str(path))

    async with ctx.activity(f"composing {name}"):
        await ctx.tool("docx.compose", {"file": name}, lambda: asyncio.to_thread(make), lambda _: f"{name} written", simulated=False)
    art = artifact_service.register(ctx.run_id, ctx.task.id, path, "docx", ctx.simulated)
    return {"artifact_id": art["id"], "filename": name, "text": text[:2000]}


IMPLS = {"fs.list": t_fs_list, "fs.search": t_fs_search, "fs.read": t_fs_read, "fs.import": t_fs_import, "fs.write_text": t_fs_write_text,
         "fs.copy": t_fs_copy, "fs.move": t_fs_move, "doc.read": t_doc_read, "docx.compose": t_docx_compose}


# ------------------------------------------------------------------ verifiers

def v_fs_list(a, out, ctx) -> list[Check]:
    fresh = sbx.list_dir(a.path)
    names = {e["name"] for e in fresh["entries"]}
    return [_ck("fs.listing_matches", {e["name"] for e in out["entries"]} == names, f"{out['count']} entries agree with a fresh scan of '{out['path']}'")]


def v_fs_search(a, out, ctx) -> list[Check]:
    bad = [m["path"] for m in out["matches"] if not sbx.resolve(m["path"], must_exist=False).exists()]
    return [_ck("fs.matches_exist", not bad, f"all {out['count']} matches exist inside the approved folders" if not bad else f"missing: {bad[:3]}")]


def v_fs_read(a, out, ctx) -> list[Check]:
    p = sbx.resolve(a.path)
    return [_ck("fs.unchanged", sha256_file(p) == out["sha256"], "file content is identical to what was read"),
            _ck("fs.has_text", out["chars"] > 0, f"{out['chars']} characters extracted" if out["chars"] else "no text could be extracted from this file")]


def v_fs_import(a, out, ctx) -> list[Check]:
    path, art = artifact_service.resolve_download(out["artifact_id"])
    src = sbx.resolve(a.path)
    return [_ck("import.hash_equal", sha256_file(path) == sha256_file(src) == out["sha256"], "the imported copy is byte-identical to the original"),
            _ck("import.registered", art.run_id == ctx.run_id, f"{art.filename} registered as a deliverable of this run")]


def v_fs_write_text(a, out, ctx) -> list[Check]:
    path, art = artifact_service.resolve_download(out["artifact_id"])
    got = path.read_text(encoding="utf-8")
    return [_ck("write.readback", _sha_text(got) == out["expected_sha256"], f"{art.filename} read back and equals the requested content"),
            _ck("write.nonempty", bool(got.strip()), f"{len(got)} characters")]


def v_fs_copy(a, out, ctx) -> list[Check]:
    src, dst = sbx.resolve(a.source), sbx.resolve(a.destination)
    return [_ck("copy.dest_exists", dst.is_file(), f"{out['destination']} exists"),
            _ck("copy.hash_equal", sha256_file(dst) == sha256_file(src) == out["sha256"], "destination is byte-identical to the source"),
            _ck("copy.source_intact", src.is_file(), "the source file is still in place")]


def v_fs_move(a, out, ctx) -> list[Check]:
    dst = sbx.resolve(a.destination)
    src_gone = not sbx.resolve(a.source, must_exist=False).exists()
    return [_ck("move.source_gone", src_gone, f"{out['source']} no longer exists" if src_gone else f"{out['source']} still exists"),
            _ck("move.hash_equal", dst.is_file() and sha256_file(dst) == out["sha256"], "destination content equals the file that was moved")]


def v_doc_read(a, out, ctx) -> list[Check]:
    from .custom_agent import find_upload

    path, _ = find_upload(ctx, a.file)
    again = tfill.read_xlsx_text(path) if path.suffix.lower() == ".xlsx" else extract_text(path)["text"]
    return [_ck("doc.reread_equal", len(again) == out["chars"], f"a fresh read gives the same {out['chars']} characters"),
            _ck("doc.has_text", out["chars"] > 0, "text extracted" if out["chars"] else "the document has no extractable text")]


def v_docx_compose(a, out, ctx) -> list[Check]:
    from .custom_agent import _rows, fmt_value

    path, art = artifact_service.resolve_download(out["artifact_id"])
    text = tfill.read_docx_text(path)
    nt = norm(text)
    first = norm(out["text"].split("\n")[0].lstrip("#- "))[:40] if out["text"].strip() else ""
    ck = [_ck("docx.readable", bool(text.strip()), f"{art.filename} reopened ({len(text)} chars)"),
          _ck("docx.title", norm(a.title)[:30] in nt, "title present"),
          _ck("docx.body", not first or first in nt, "body text present")]
    if a.table_source:
        rows = _rows(resolve_args(a.table_source, ctx.outputs))
        missing = [fmt_value(v) for r in rows[:5] for v in r.values() if fmt_value(v) and norm(fmt_value(v)) not in nt]
        ck.append(_ck("docx.table", not missing, "table values present" if not missing else f"missing table values: {missing[:3]}"))
    return ck


VERIFIERS = {"fs.list": v_fs_list, "fs.search": v_fs_search, "fs.read": v_fs_read, "fs.import": v_fs_import, "fs.write_text": v_fs_write_text,
             "fs.copy": v_fs_copy, "fs.move": v_fs_move, "doc.read": v_doc_read, "docx.compose": v_docx_compose}
