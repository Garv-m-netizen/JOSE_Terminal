"""Excel workbook generation (openpyxl) and independent read-back (inspection)."""
from __future__ import annotations

from datetime import datetime, timezone
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from ..schemas.domain import CompetitorRecord
from .report_stats import SIM_BANNER, STATUS_LABEL, compute_stats

COMPARISON_HEADERS = ["Competitor", "Product", "Website", "Summary", "Key features", "Pricing status",
                      "Starting price (USD)", "Pricing details", "Verification", "Source URLs",
                      "Retrieved at (UTC)", "Data type"]
SOURCES_HEADERS = ["Evidence ID", "Competitor", "Claim type", "Claim", "Source URL", "Page title",
                   "Retrieved at (UTC)", "Evidence excerpt", "Status"]
REQUIRED_SHEETS = ["Summary", "Comparison", "Sources"]
HEADER_FILL = PatternFill("solid", fgColor="1F2A44")
HEADER_FONT = Font(bold=True, color="FFFFFF")
THIN = Side(style="thin", color="D0D5DD")


def _style_header(ws, ncols: int) -> None:
    for c in range(1, ncols + 1):
        cell = ws.cell(row=1, column=c)
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    ws.freeze_panes = "A2"


def _widths(ws, widths: list[int]) -> None:
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def _fmt_ts(dt: datetime) -> str:
    return dt.astimezone(timezone.utc).strftime("%Y-%m-%d %H:%M")


def build_workbook(path: Path, records: list[CompetitorRecord], evidence: list[dict], meta: dict) -> None:
    simulated = bool(meta.get("simulated"))
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    st = compute_stats(records)
    rows = [("Competitor comparison", meta.get("industry", "")), ("Generated (UTC)", _fmt_ts(datetime.now(timezone.utc))),
            ("Data type", SIM_BANNER if simulated else "Real research from retrieved web pages"),
            ("Competitors", st["count"]), ("With public price", st["with_public_price"]),
            ("With free tier", st["free_tier"]), ("Contact sales only", st["contact_sales"]),
            ("Pricing not available", st["not_found"]), ("Fully verified records", st["verified"]),
            ("Unique sources", st["unique_sources"]), ("Evidence records", len(evidence)),
            ("Methodology", "Pages were retrieved at run time; a claim is kept only if its quoted evidence appears verbatim in the page text. "
                            "Unavailable pricing is left blank and labelled, never recorded as zero.")]
    for r in rows:
        ws.append(list(r))
    ws["A1"].font = Font(bold=True, size=14)
    for r in range(2, len(rows) + 1):
        ws.cell(row=r, column=1).font = Font(bold=True)
    ws.cell(row=len(rows), column=2).alignment = Alignment(wrap_text=True, vertical="top")
    if simulated:
        ws["B3"].font = Font(bold=True, color="B42318")
    _widths(ws, [26, 90])

    cs = wb.create_sheet("Comparison")
    cs.append(COMPARISON_HEADERS)
    for r in records:
        plans = "; ".join(f"{p.name}: {p.price_text}" for p in r.plans) or (r.pricing_note or STATUS_LABEL[r.pricing_status])
        cs.append([
            r.name, r.product, r.website, r.summary, "\n".join(f"• {f.text}" for f in r.features),
            STATUS_LABEL[r.pricing_status], r.starting_price_usd, plans, r.verification,
            "\n".join(dict.fromkeys(s.url for s in r.sources)),
            _fmt_ts(max(s.retrieved_at for s in r.sources)) if r.sources else "",
            "SIMULATED" if r.simulated else "Real"])
        row = cs.max_row
        cs.cell(row=row, column=3).hyperlink = r.website
        cs.cell(row=row, column=3).font = Font(color="1D4ED8", underline="single")
        cs.cell(row=row, column=7).number_format = '"$"#,##0.00'
        for c in range(1, len(COMPARISON_HEADERS) + 1):
            cs.cell(row=row, column=c).alignment = Alignment(wrap_text=True, vertical="top")
            cs.cell(row=row, column=c).border = Border(bottom=THIN)
    _style_header(cs, len(COMPARISON_HEADERS))
    _widths(cs, [22, 24, 32, 42, 48, 24, 18, 44, 14, 48, 18, 12])
    cs.auto_filter.ref = f"A1:{get_column_letter(len(COMPARISON_HEADERS))}{max(cs.max_row, 2)}"

    ss = wb.create_sheet("Sources")
    ss.append(SOURCES_HEADERS)
    for e in evidence:
        ss.append([e["id"], e["subject"], e["claim_type"], e["claim"], e["source_url"], e["page_title"],
                   e["retrieved_at"], e["evidence_text"], e["status"]])
        for c in range(1, len(SOURCES_HEADERS) + 1):
            ss.cell(row=ss.max_row, column=c).alignment = Alignment(wrap_text=True, vertical="top")
    _style_header(ss, len(SOURCES_HEADERS))
    _widths(ss, [16, 22, 12, 40, 48, 32, 18, 60, 14])
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def inspect_workbook(path: Path) -> dict:
    """Independent read-back used by the Verification Agent."""
    wb = load_workbook(path, read_only=False, data_only=True)
    out: dict = {"sheets": wb.sheetnames}
    cs = wb["Comparison"] if "Comparison" in wb.sheetnames else None
    if cs is not None:
        rows = list(cs.iter_rows(values_only=True))
        header = [str(h) for h in rows[0]] if rows else []
        out["comparison_headers"] = header
        out["comparison_rows"] = [dict(zip(header, r)) for r in rows[1:] if any(v is not None for v in r)]
    ss = wb["Sources"] if "Sources" in wb.sheetnames else None
    if ss is not None:
        rows = list(ss.iter_rows(values_only=True))
        header = [str(h) for h in rows[0]] if rows else []
        out["sources_headers"] = header
        out["sources_rows"] = [dict(zip(header, r)) for r in rows[1:] if any(v is not None for v in r)]
    sm = wb["Summary"] if "Summary" in wb.sheetnames else None
    if sm is not None:
        out["summary"] = {str(r[0]): r[1] for r in sm.iter_rows(values_only=True) if r[0]}
    wb.close()
    return out
