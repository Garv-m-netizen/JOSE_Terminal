"""Controlled file access: validated uploads, safe reads, hashing."""
from __future__ import annotations

import hashlib
from pathlib import Path

ALLOWED_UPLOAD_EXT = {".pdf", ".txt", ".md", ".csv", ".xlsx", ".docx", ".png", ".jpg", ".jpeg"}
MAX_EXCERPT_CHARS = 6000


class FileRejected(ValueError):
    pass


def validate_upload(filename: str, size: int, max_mb: int) -> str:
    name = Path(filename).name  # strip any directory component
    ext = Path(name).suffix.lower()
    if ext not in ALLOWED_UPLOAD_EXT:
        raise FileRejected(f"file type '{ext or '(none)'}' not allowed; allowed: {', '.join(sorted(ALLOWED_UPLOAD_EXT))}")
    if size <= 0:
        raise FileRejected("file is empty")
    if size > max_mb * 1024 * 1024:
        raise FileRejected(f"file exceeds the {max_mb} MB limit")
    return name


def sha256_file(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()


def extract_text(path: Path) -> dict:
    """Read a supported input document -> {kind, text, pages/rows}. Raises FileRejected on unreadable input."""
    ext = path.suffix.lower()
    try:
        if ext == ".pdf":
            import pymupdf as fitz

            with fitz.open(path) as doc:
                pages = [p.get_text() for p in doc]
            text = "\n".join(pages).strip()
            # OCR is only needed for scanned PDFs; we report that honestly instead of guessing.
            return {"kind": "pdf", "text": text, "pages": len(pages),
                    "needs_ocr": not text}
        if ext in (".txt", ".md"):
            return {"kind": "text", "text": path.read_text(encoding="utf-8", errors="replace")}
        if ext == ".docx":
            import docx

            d = docx.Document(str(path))
            parts = [p.text for p in d.paragraphs if p.text.strip()]
            for t in d.tables:
                parts += [" | ".join(c.text.strip() for c in r.cells) for r in t.rows]
            return {"kind": "docx", "text": chr(10).join(parts)}
        if ext in (".png", ".jpg", ".jpeg"):
            from PIL import Image

            with Image.open(path) as im:
                return {"kind": "image", "text": "", "size": list(im.size), "needs_ocr": True}
        if ext == ".csv":
            import pandas as pd

            df = pd.read_csv(path)
            return {"kind": "table", "text": df.head(200).to_csv(index=False), "rows": len(df), "columns": list(df.columns)}
        if ext == ".xlsx":
            import pandas as pd

            sheets = pd.read_excel(path, sheet_name=None)
            parts = [f"## {n}\n{df.head(100).to_csv(index=False)}" for n, df in sheets.items()]
            return {"kind": "table", "text": "\n".join(parts), "sheets": list(sheets)}
    except FileRejected:
        raise
    except Exception as e:  # noqa: BLE001
        raise FileRejected(f"could not read {path.name}: {e}") from e
    raise FileRejected(f"unsupported file type {ext}")
