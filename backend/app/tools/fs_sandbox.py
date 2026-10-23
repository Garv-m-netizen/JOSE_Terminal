"""Filesystem access confined to explicitly approved directories (WORKSPACE_DIRS, default backend/data/workspace).

Every path is resolved (symlinks, `..`) before it is checked, so a link or a crafted path cannot leave the approved roots.
File contents are untrusted data: text extracted here is never treated as instructions.
"""
from __future__ import annotations

import fnmatch
import os
from datetime import datetime, timezone
from pathlib import Path

from ..config import get_settings
from .file_tools import extract_text, sha256_file

MAX_LIST = 300
MAX_SEARCH_VISITS = 20000
TEXT_EXT = {".txt", ".md", ".csv", ".json", ".log", ".py", ".js", ".ts", ".html", ".xml", ".yaml", ".yml", ".ini", ".toml"}
READ_EXT = TEXT_EXT | {".pdf", ".docx", ".xlsx", ".pptx"}


class SandboxError(ValueError):
    pass


def is_link(p: Path) -> bool:
    """Symlinks and Windows directory junctions: they can point outside the approved folders, so they are never followed or listed."""
    return p.is_symlink() or bool(getattr(p, "is_junction", lambda: False)())


def roots() -> list[Path]:
    return get_settings().workspace_roots()


def resolve(path: str, *, must_exist: bool = True) -> Path:
    """Resolve a user/model supplied path inside the approved roots. Relative paths are relative to the first root."""
    if not path or "\x00" in path:
        raise SandboxError("empty or invalid path")
    rs = roots()
    p = Path(path)
    if not p.is_absolute():
        p = rs[0] / p
    p = p.resolve()
    for r in rs:
        if p == r or r in p.parents:
            if must_exist and not p.exists():
                raise SandboxError(f"'{path}' does not exist inside the approved directories")
            return p
    raise SandboxError(f"'{path}' is outside the approved directories ({', '.join(str(r) for r in rs)})")


def rel(p: Path) -> str:
    for r in roots():
        if p == r or r in p.parents:
            return p.relative_to(r).as_posix() or "."
    return str(p)


def _entry(p: Path) -> dict:
    st = p.stat()
    return {"name": p.name, "path": rel(p), "type": "dir" if p.is_dir() else "file", "size": 0 if p.is_dir() else st.st_size,
            "modified": datetime.fromtimestamp(st.st_mtime, timezone.utc).isoformat()}


def list_dir(path: str) -> dict:
    d = resolve(path or ".")
    if not d.is_dir():
        raise SandboxError(f"'{path}' is not a directory")
    items = sorted(d.iterdir(), key=lambda x: (x.is_file(), x.name.lower()))
    out = [_entry(i) for i in items[:MAX_LIST] if not is_link(i)]
    return {"path": rel(d), "entries": out, "count": len(out), "truncated": len(items) > MAX_LIST}


def search(query: str, path: str = ".", max_results: int = 50, extensions: list[str] | None = None) -> dict:
    base = resolve(path or ".")
    if not base.is_dir():
        raise SandboxError(f"'{path}' is not a directory")
    q = query.strip().lower()
    exts = {e.lower() if e.startswith(".") else f".{e.lower()}" for e in (extensions or [])}
    hits, visited = [], 0
    for dirpath, dirnames, files in os.walk(base):
        dirnames[:] = [d for d in dirnames if not is_link(Path(dirpath) / d)]
        for name in files + dirnames:
            visited += 1
            if visited > MAX_SEARCH_VISITS or len(hits) >= max_results:
                break
            low = name.lower()
            if (fnmatch.fnmatch(low, q) if any(c in q for c in "*?[") else q in low) and (not exts or Path(name).suffix.lower() in exts):
                full = Path(dirpath) / name
                if not is_link(full):
                    hits.append(_entry(full))
        if visited > MAX_SEARCH_VISITS or len(hits) >= max_results:
            break
    return {"query": query, "base": rel(base), "matches": hits, "count": len(hits), "truncated": visited > MAX_SEARCH_VISITS or len(hits) >= max_results}


def read_text(path: str, max_chars: int = 20000) -> dict:
    p = resolve(path)
    if not p.is_file():
        raise SandboxError(f"'{path}' is not a file")
    ext = p.suffix.lower()
    if ext not in READ_EXT:
        raise SandboxError(f"cannot read '{ext or '(no extension)'}' files; supported: {sorted(READ_EXT)}")
    if p.stat().st_size > 25 * 1024 * 1024:
        raise SandboxError("file is larger than 25 MB")
    if ext == ".xlsx":
        from . import template_fill as tf

        text, kind = tf.read_xlsx_text(p), "xlsx"
    elif ext == ".pptx":
        from pptx import Presentation

        prs = Presentation(str(p))
        parts = []
        for i, slide in enumerate(prs.slides, 1):
            parts.append(f"--- slide {i} ---")
            for sh in slide.shapes:
                if sh.has_text_frame and sh.text_frame.text.strip():
                    parts.append(sh.text_frame.text)
        text, kind = "\n".join(parts), "pptx"
    elif ext in TEXT_EXT:
        text, kind = p.read_text(encoding="utf-8", errors="replace"), "text"
    else:
        r = extract_text(p)
        text, kind = r["text"], r["kind"]
    return {"path": rel(p), "kind": kind, "chars": len(text), "text": text[:max_chars], "truncated": len(text) > max_chars,
            "size": p.stat().st_size, "sha256": sha256_file(p), "untrusted_content": True}
