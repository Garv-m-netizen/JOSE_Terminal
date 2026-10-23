"""PowerPoint generation (python-pptx) and independent read-back."""
from __future__ import annotations

from datetime import datetime, timezone
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Emu, Inches, Pt

from ..schemas.domain import CompetitorRecord
from .report_stats import SIM_BANNER, STATUS_LABEL, compute_stats, key_findings, limitations

NAVY, BLUE, GREY, RED = RGBColor(0x14, 0x1B, 0x2D), RGBColor(0x1D, 0x4E, 0xD8), RGBColor(0x55, 0x5F, 0x73), RGBColor(0xB4, 0x23, 0x18)
REQUIRED_TITLES = ["Methodology", "Competitor comparison", "Key findings", "Limitations", "References"]
ROWS_PER_TABLE = 5
REFS_PER_SLIDE = 9


def _slide(prs, title: str, simulated: bool):
    s = prs.slides.add_slide(prs.slide_layouts[5])  # Title Only
    t = s.shapes.title
    t.text = title
    t.left, t.top, t.width, t.height = Inches(0.6), Inches(0.35), Inches(12.1), Inches(0.9)
    for p in t.text_frame.paragraphs:
        for r in p.runs:
            r.font.size, r.font.bold, r.font.color.rgb = Pt(30), True, NAVY
    if simulated:
        _text(s, SIM_BANNER, 0.6, 7.0, 12.1, 0.35, 11, RED, bold=True)
    return s


def _text(slide, text: str, x: float, y: float, w: float, h: float, size: int = 16, color=NAVY, bold: bool = False):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.text = text
    for p in tf.paragraphs:
        for r in p.runs:
            r.font.size, r.font.color.rgb, r.font.bold = Pt(size), color, bold
    return tb


def _bullets(slide, items: list[str], x=0.7, y=1.5, w=11.9, h=5.2, size=17):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    for i, it in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = f"•  {it}"
        p.space_after = Pt(10)
        for r in p.runs:
            r.font.size, r.font.color.rgb = Pt(size), NAVY


def build_presentation(path: Path, records: list[CompetitorRecord], meta: dict) -> None:
    simulated, industry = bool(meta.get("simulated")), meta.get("industry", "")
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)

    s = prs.slides.add_slide(prs.slide_layouts[5])
    s.shapes.title.text = f"Competitor landscape: {industry}"
    s.shapes.title.left, s.shapes.title.top, s.shapes.title.width, s.shapes.title.height = Inches(0.8), Inches(2.2), Inches(11.7), Inches(1.6)
    for r in s.shapes.title.text_frame.paragraphs[0].runs:
        r.font.size, r.font.bold, r.font.color.rgb = Pt(40), True, NAVY
    _text(s, f"{len(records)} competitors compared on features and public pricing\nGenerated {datetime.now(timezone.utc):%Y-%m-%d} (UTC) by JOSE",
          0.8, 4.0, 11.7, 1.2, 20, GREY)
    if simulated:
        _text(s, SIM_BANNER, 0.8, 5.6, 11.7, 0.6, 18, RED, bold=True)

    st = compute_stats(records)
    s = _slide(prs, "Methodology", simulated)
    _bullets(s, [f"Identified {st['count']} competitors in {industry} via web search.",
                 "Retrieved each competitor's homepage, pricing page and one independent page at run time.",
                 "Kept a feature or price only if its quoted evidence appears verbatim in the retrieved page text.",
                 "Recorded unavailable pricing as unavailable - never as zero.",
                 f"{st['total_features']} evidence-backed feature claims from {st['unique_sources']} distinct sources.",
                 "Every generated file was reopened and checked against the stored evidence before delivery."])

    chunks = [records[i:i + ROWS_PER_TABLE] for i in range(0, len(records), ROWS_PER_TABLE)] or [[]]
    for ci, chunk in enumerate(chunks):
        title = "Competitor comparison" + (f" ({ci + 1}/{len(chunks)})" if len(chunks) > 1 else "")
        s = _slide(prs, title, simulated)
        shape = s.shapes.add_table(len(chunk) + 1, 5, Inches(0.6), Inches(1.5), Inches(12.1), Inches(0.5 * (len(chunk) + 1)))
        tbl = shape.table
        for i, w in enumerate([2.2, 2.4, 2.6, 1.6, 3.3]):
            tbl.columns[i].width = Inches(w)
        for i, h in enumerate(["Competitor", "Product", "Pricing", "From (USD)", "Top features"]):
            tbl.cell(0, i).text = h
        for ri, r in enumerate(chunk, 1):
            price = f"${r.starting_price_usd:,.0f}" if r.starting_price_usd is not None else "n/a"
            vals = [r.name, r.product, STATUS_LABEL[r.pricing_status], price, "; ".join(f.text for f in r.features[:3])]
            for ci2, v in enumerate(vals):
                tbl.cell(ri, ci2).text = v
        for row in tbl.rows:
            row.height = Emu(int(Inches(0.9)))
            for cell in row.cells:
                for p in cell.text_frame.paragraphs:
                    for run in p.runs:
                        run.font.size = Pt(12)

    s = _slide(prs, "Key findings", simulated)
    _bullets(s, key_findings(records))
    s = _slide(prs, "Limitations", simulated)
    _bullets(s, limitations(records, simulated), size=16)

    refs: list[str] = []
    for r in records:
        for src in dict.fromkeys(x.url for x in r.sources):
            refs.append(f"{r.name}: {src}")
    for inp in meta.get("input_files", []):
        refs.append(f"Input document: {inp}")
    pages = [refs[i:i + REFS_PER_SLIDE] for i in range(0, len(refs), REFS_PER_SLIDE)] or [["No sources recorded."]]
    for pi, page in enumerate(pages):
        s = _slide(prs, "References" + (f" ({pi + 1}/{len(pages)})" if len(pages) > 1 else ""), simulated)
        _bullets(s, page, size=13)
    path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(path)


def inspect_presentation(path: Path) -> dict:
    prs = Presentation(path)
    slides = []
    for sl in prs.slides:
        title = sl.shapes.title.text if sl.shapes.title is not None else ""
        texts = []
        for sh in sl.shapes:
            if sh.has_text_frame:
                texts.append(sh.text_frame.text)
            if getattr(sh, "has_table", False) and sh.has_table:
                for row in sh.table.rows:
                    texts.append(" | ".join(c.text for c in row.cells))
        slides.append({"title": title, "text": "\n".join(texts)})
    return {"slide_count": len(slides), "slides": slides}
