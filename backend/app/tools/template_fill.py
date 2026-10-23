"""Fill Word / PDF / Excel templates from a data row, and convert Office files to PDF (LibreOffice or MS Office via COM).

Mapping is {template field -> value}. A template field can be written in any of these styles (all matched case/punctuation-insensitively):
  * placeholder keys:        {{Name}}  <<Name>>  [[Name]]  ${Name}  {Name}
  * label + bracket:         "Student Name: [STUDENT NAME]"   -> key "Student Name" (or "Student Name:")
  * bare bracket text:       "[STUDENT NAME]"                  -> key "STUDENT NAME"
  * label + blank:           "Father's Name: ______"  or an empty cell next to the label
Every key is reported as filled or not, so callers can verify.
"""
from __future__ import annotations

import os
import re
import shutil
import subprocess
import sys
import threading
from pathlib import Path

from .inspect_files import BRACKET_RE, LABELED_BRACKET_RE, PLACEHOLDER_RE, norm_key

_WORD_LOCK = threading.Lock()
BLANK_LABEL_RE = re.compile(r"(?:^|[|\s])([A-Za-z][A-Za-z0-9 .'/&()\-]{1,45}?)\s*:\s*(_{2,}|\.{3,})?\s*$")


class FillError(RuntimeError):
    pass


def _norm_values(mapping: dict[str, str]) -> dict[str, str]:
    return {norm_key(k): v for k, v in mapping.items()}


def find_spans(text: str, values: dict[str, str], hit: set[str]) -> list[tuple[int, int, str]]:
    """Locate every place in `text` that a mapped field should be written: [(start, end, replacement)] (non-overlapping)."""
    spans: list[tuple[int, int, str]] = []

    def free(a: int, b: int) -> bool:
        return all(b <= s or a >= e for s, e, _ in spans) if a != b else all(not (s < a < e) for s, e, _ in spans)

    def add(a: int, b: int, n: str, v: str) -> None:
        if free(a, b):
            spans.append((a, b, v))
            hit.add(n)

    for m in PLACEHOLDER_RE.finditer(text):  # {{Name}} <<Name>> [[Name]] ${Name} {Name}
        n = norm_key(next(g for g in m.groups() if g))
        if n in values:
            add(m.start(), m.end(), n, values[n])
    for m in LABELED_BRACKET_RE.finditer(text):  # Student Name: [STUDENT NAME]
        n = norm_key(m.group(1))
        if n in values:
            add(m.start(2), m.end(2), n, values[n])
    for m in BRACKET_RE.finditer(text):  # [STUDENT NAME]
        n = norm_key(m.group(1))
        if n in values:
            add(m.start(), m.end(), n, values[n])
    m = BLANK_LABEL_RE.search(text)  # Father's Name: ____   (or nothing after the colon)
    if m and norm_key(m.group(1)) in values:
        n = norm_key(m.group(1))
        if m.group(2):
            add(m.start(2), m.end(2), n, values[n])
        else:
            add(len(text.rstrip()), len(text.rstrip()), n, " " + values[n])
    return spans


def _sub_text(text: str, values: dict[str, str], hit: set[str]) -> str:
    out = text
    for a, b, v in sorted(find_spans(text, values, hit), reverse=True):
        out = out[:a] + v + out[b:]
    return out


# ------------------------------------------------------------------ DOCX

def _apply_spans_to_runs(p, spans: list[tuple[int, int, str]]) -> None:
    """Write replacements into the paragraph's runs so the surrounding formatting (bold labels etc.) is preserved."""
    for a, b, new in sorted(spans, reverse=True):
        runs = p.runs
        pos, placed = 0, False
        for r in runs:
            s, e = pos, pos + len(r.text)
            pos = e
            if a == b:  # pure insertion after text that ends at a
                if s < a <= e and not placed:
                    r.text = r.text[: a - s] + new + r.text[a - s:]
                    placed = True
                continue
            if e <= a or s >= b:
                continue
            lo, hi = max(a, s) - s, min(b, e) - s
            r.text = r.text[:lo] + (new if not placed else "") + r.text[hi:]
            placed = True


def _fill_paragraph(p, values: dict[str, str], hit: set[str]) -> None:
    full = "".join(r.text for r in p.runs)
    if not full.strip():
        return
    spans = find_spans(full, values, hit)
    if spans:
        _apply_spans_to_runs(p, spans)


def _set_cell(cell, value: str) -> None:
    p = cell.paragraphs[0] if cell.paragraphs else None
    if p is None:
        cell.text = value
        return
    if p.runs:
        p.runs[0].text = value
        for r in p.runs[1:]:
            r.text = ""
    else:
        p.add_run(value)


def _fill_grid(table, values: dict[str, str], hit: set[str]) -> None:
    """Row-label x column-header addressing for table templates: key 'Maths / Marks Obtained'."""
    rows = table.rows
    if len(rows) < 2:
        return
    header = [c.text.strip() for c in rows[0].cells]
    if sum(1 for h in header if h) < 2:
        return
    for row in rows[1:]:
        cells = row.cells
        label = cells[0].text.strip()
        if not label:
            continue
        for j in range(1, len(cells)):
            if not header[j] or cells[j]._tc is cells[j - 1]._tc:
                continue
            n = norm_key(label + " " + header[j])
            if n in values:
                _set_cell(cells[j], values[n])
                hit.add(n)


def fill_docx(template: Path, out: Path, mapping: dict[str, str]) -> dict:
    import docx

    d = docx.Document(str(template))
    values, hit = _norm_values(mapping), set()

    def walk_tables(tables) -> None:
        for t in tables:
            _fill_grid(t, values, hit)
            for row in t.rows:
                cells = row.cells
                for i, c in enumerate(cells):
                    for p in c.paragraphs:
                        _fill_paragraph(p, values, hit)
                    # label alone in this cell, empty value cell next to it
                    label = c.text.strip().rstrip(":").strip()
                    if label and norm_key(label) in values and i + 1 < len(cells) and not cells[i + 1].text.strip() and cells[i + 1]._tc is not c._tc:
                        n = norm_key(label)
                        nxt = cells[i + 1]
                        (nxt.paragraphs[0].add_run(values[n]) if nxt.paragraphs else setattr(nxt, "text", values[n]))
                        hit.add(n)
                    walk_tables(c.tables)

    for p in d.paragraphs:
        _fill_paragraph(p, values, hit)
    walk_tables(d.tables)
    for sec in d.sections:
        for part in (sec.header, sec.footer):
            for p in part.paragraphs:
                _fill_paragraph(p, values, hit)
            walk_tables(part.tables)
    out.parent.mkdir(parents=True, exist_ok=True)
    d.save(str(out))
    return {"filled": sorted(k for k in mapping if norm_key(k) in hit), "unfilled": sorted(k for k in mapping if norm_key(k) not in hit)}


def read_docx_text(path: Path) -> str:
    import docx

    d = docx.Document(str(path))
    parts = [p.text for p in d.paragraphs]

    def tables(ts) -> None:
        for t in ts:
            for r in t.rows:
                for c in r.cells:
                    parts.append(c.text)
                    tables(c.tables)

    tables(d.tables)
    for sec in d.sections:
        parts += [p.text for p in sec.header.paragraphs + sec.footer.paragraphs]
    return "\n".join(parts)


# ------------------------------------------------------------------ XLSX

def fill_xlsx(template: Path, out: Path, mapping: dict[str, str]) -> dict:
    import openpyxl

    wb = openpyxl.load_workbook(template)
    values, hit = _norm_values(mapping), set()
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if not isinstance(c.value, str):
                    continue
                new = _sub_text(c.value, values, hit)
                if new != c.value:
                    c.value = new
                    continue
                label = c.value.strip().rstrip(":").strip()
                n = norm_key(label)
                if n in values and c.value.strip().endswith(":"):
                    right = ws.cell(row=c.row, column=c.column + 1)
                    if right.value in (None, ""):
                        right.value = values[n]
                        hit.add(n)
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    return {"filled": sorted(k for k in mapping if norm_key(k) in hit), "unfilled": sorted(k for k in mapping if norm_key(k) not in hit)}


def read_xlsx_text(path: Path) -> str:
    import openpyxl

    wb = openpyxl.load_workbook(path, data_only=False)
    return "\n".join(str(c.value) for ws in wb.worksheets for row in ws.iter_rows() for c in row if c.value is not None)


# ------------------------------------------------------------------ PDF

def _font_for(value: str):
    import pymupdf

    if any(ord(ch) > 255 for ch in value):
        for f in (r"C:\Windows\Fonts\Nirmala.ttf", r"C:\Windows\Fonts\arial.ttf", "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"):
            if os.path.exists(f):
                return pymupdf.Font(fontfile=f), f
    return pymupdf.Font("helv"), None


def fill_pdf(template: Path, out: Path, mapping: dict[str, str]) -> dict:
    import pymupdf

    values = _norm_values(mapping)
    raw_keys = {norm_key(k): k for k in mapping}
    hit: set[str] = set()
    doc = pymupdf.open(template)
    for page in doc:
        for w in page.widgets() or []:  # real form fields
            n = norm_key(w.field_name or "")
            if n in values:
                w.field_value = values[n]
                w.update()
                hit.add(n)
        text = page.get_text()
        todo: list[tuple[pymupdf.Rect, str]] = []  # rect to white-out and the value to write there

        def rects_of(token: str, near: pymupdf.Rect | None = None) -> list[pymupdf.Rect]:
            found = page.search_for(token)
            if near is not None and found:  # the bracket that sits on the same line, to the right of the label
                found = [r for r in found if abs(r.y0 - near.y0) < 6 and r.x0 >= near.x0] or found
                found = sorted(found, key=lambda r: (abs(r.y0 - near.y0), r.x0))[:1]
            return found

        for m in PLACEHOLDER_RE.finditer(text):
            n = norm_key(next(g for g in m.groups() if g))
            if n in values:
                todo += [(r, values[n]) for r in rects_of(m.group(0))]
                hit.add(n)
        for m in LABELED_BRACKET_RE.finditer(text):
            n = norm_key(m.group(1))
            if n in values and n not in hit:
                label = rects_of(m.group(1).strip() + ":")
                got = rects_of(m.group(2), label[0] if label else None)
                if got:
                    todo += [(got[0], values[n])]
                    hit.add(n)
        for m in BRACKET_RE.finditer(text):
            n = norm_key(m.group(1))
            if n in values and n not in hit:
                rs = rects_of(m.group(0))
                if rs:
                    todo += [(r, values[n]) for r in rs]
                    hit.add(n)
        for r, _ in todo:
            page.add_redact_annot(r, fill=False)  # remove only the text, never paint over borders
        if todo:
            page.apply_redactions(images=0, graphics=0)
        for r, v in todo:
            _, ff = _font_for(v)
            size = max(7, min(12, r.height * 0.85))
            kw = {"fontfile": ff, "fontname": "F0"} if ff else {"fontname": "helv"}
            page.insert_text((r.x0, r.y1 - 2), v, fontsize=size, **kw)
        for n, v in values.items():  # bare labels: write right after "Label:"
            if n in hit:
                continue
            rs = page.search_for(raw_keys[n].strip().rstrip(":") + ":")
            if rs:
                r = rs[0]
                _, ff = _font_for(v)
                kw = {"fontfile": ff, "fontname": "F0"} if ff else {"fontname": "helv"}
                page.insert_text((r.x1 + 6, r.y1 - 2.5), v, fontsize=max(8, min(12, r.height * 0.85)), **kw)
                hit.add(n)
    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out, garbage=3, deflate=True)
    doc.close()
    return {"filled": sorted(k for k in mapping if norm_key(k) in hit), "unfilled": sorted(k for k in mapping if norm_key(k) not in hit)}


def read_pdf_text(path: Path) -> str:
    import pymupdf

    with pymupdf.open(path) as doc:
        return "\n".join(p.get_text() for p in doc)


# ------------------------------------------------------------------ PDF conversion

def find_converter() -> str | None:
    """Return 'libreoffice', 'word' or None. Only fixed, built-in commands are ever executed."""
    if _soffice():
        return "libreoffice"
    if sys.platform == "win32":
        try:
            import win32com.client  # noqa: F401

            for p in (r"C:\Program Files\Microsoft Office\root\Office16\WINWORD.EXE", r"C:\Program Files (x86)\Microsoft Office\root\Office16\WINWORD.EXE",
                      r"C:\Program Files\Microsoft Office\Office16\WINWORD.EXE"):
                if os.path.exists(p):
                    return "word"
        except ImportError:
            return None
    return None


def _soffice() -> str | None:
    found = shutil.which("soffice") or shutil.which("libreoffice")
    if found:
        return found
    for p in (r"C:\Program Files\LibreOffice\program\soffice.exe", r"C:\Program Files (x86)\LibreOffice\program\soffice.exe"):
        if os.path.exists(p):
            return p
    return None


def convert_to_pdf(sources: list[Path], out_dir: Path) -> dict[Path, Path]:
    """Convert docx/xlsx files to PDF. Returns {source: pdf}. Raises FillError when no converter exists."""
    out_dir.mkdir(parents=True, exist_ok=True)
    conv = find_converter()
    if conv is None:
        raise FillError("No PDF converter found. Install LibreOffice (or use Windows with Microsoft Word) to produce PDFs from Word/Excel templates, "
                        "or ask for Word output instead.")
    result: dict[Path, Path] = {}
    if conv == "libreoffice":
        exe = _soffice()
        for s in sources:
            r = subprocess.run([exe, "--headless", "--convert-to", "pdf", "--outdir", str(out_dir), str(s)], capture_output=True, timeout=180)  # noqa: S603
            pdf = out_dir / (s.stem + ".pdf")
            if r.returncode != 0 or not pdf.exists():
                raise FillError(f"LibreOffice failed to convert {s.name}")
            result[s] = pdf
        return result
    import pythoncom
    import win32com.client

    with _WORD_LOCK:
        pythoncom.CoInitialize()
        word = excel = None
        try:
            for s in sources:
                pdf = out_dir / (s.stem + ".pdf")
                if s.suffix.lower() == ".docx":
                    word = word or win32com.client.DispatchEx("Word.Application")
                    word.Visible, word.DisplayAlerts = False, 0
                    doc = word.Documents.Open(str(s.resolve()), ReadOnly=True, AddToRecentFiles=False)
                    try:
                        doc.ExportAsFixedFormat(str(pdf.resolve()), 17)  # wdExportFormatPDF
                    finally:
                        doc.Close(False)
                elif s.suffix.lower() == ".xlsx":
                    excel = excel or win32com.client.DispatchEx("Excel.Application")
                    excel.Visible, excel.DisplayAlerts = False, False
                    wb = excel.Workbooks.Open(str(s.resolve()), ReadOnly=True)
                    try:
                        wb.ExportAsFixedFormat(0, str(pdf.resolve()))  # xlTypePDF
                    finally:
                        wb.Close(False)
                else:
                    raise FillError(f"cannot convert {s.suffix} to PDF")
                if not pdf.exists():
                    raise FillError(f"conversion of {s.name} produced no PDF")
                result[s] = pdf
        except FillError:
            raise
        except Exception as e:  # noqa: BLE001
            raise FillError(f"Microsoft Office could not convert the file: {e}") from e
        finally:
            for app in (word, excel):
                try:
                    if app is not None:
                        app.Quit()
                except Exception:  # noqa: BLE001
                    pass
            # CoUninitialize is deliberately skipped: the worker thread is reused, and uninitialising while wrappers are
            # still alive makes Windows report an RPC fault. Leaving COM initialised on a pool thread is harmless.
    return result
