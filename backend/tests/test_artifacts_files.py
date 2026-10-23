from __future__ import annotations

import io
from datetime import datetime, timezone

import pytest
from openpyxl import load_workbook

from app.database import session_scope
from app.models import Artifact
from app.schemas.domain import CompetitorRecord, Feature, Plan, SourceRef
from app.tools import excel_tools, powerpoint_tools
from app.tools.file_tools import FileRejected, extract_text, validate_upload
from app.tools.safety import UnsafePath, safe_join

from .conftest import OBJECTIVE, wait_status


def records(simulated=False):
    now = datetime(2026, 1, 1, tzinfo=timezone.utc)
    out = []
    for i, st in enumerate(["paid", "contact_sales", "not_found"]):
        out.append(CompetitorRecord(
            name=f"Co{i}", product=f"Co{i} Suite", website=f"https://co{i}.test", summary="s",
            features=[Feature(text=f"feat{i}", evidence_id=f"ev_f{i}")], pricing_status=st,
            plans=[Plan(name="Pro", price_text="$10 per month", price_usd=10.0, evidence_id=f"ev_p{i}")] if st == "paid" else [],
            sources=[SourceRef(url=f"https://co{i}.test", title="t", retrieved_at=now)], verification="verified", simulated=simulated))
    return out


def test_excel_generation_and_readback(tmp_path):
    p = tmp_path / "c.xlsx"
    ev = [{"id": "ev_f0", "subject": "Co0", "claim_type": "feature", "claim": "feat0", "source_url": "https://co0.test", "page_title": "t",
           "retrieved_at": "2026-01-01 00:00", "evidence_text": "feat0 text", "status": "verified"}]
    excel_tools.build_workbook(p, records(), ev, {"industry": "X", "simulated": False})
    info = excel_tools.inspect_workbook(p)
    assert info["sheets"] == ["Summary", "Comparison", "Sources"]
    assert info["comparison_headers"] == excel_tools.COMPARISON_HEADERS
    rows = info["comparison_rows"]
    assert [r["Competitor"] for r in rows] == ["Co0", "Co1", "Co2"]
    assert rows[0]["Starting price (USD)"] == 10 and rows[1]["Starting price (USD)"] is None  # unavailable != 0
    assert rows[2]["Pricing status"] == "Not available"
    assert len(info["sources_rows"]) == 1


def test_pptx_generation_and_readback(tmp_path):
    p = tmp_path / "c.pptx"
    powerpoint_tools.build_presentation(p, records(simulated=True), {"industry": "X", "simulated": True})
    info = powerpoint_tools.inspect_presentation(p)
    titles = [s["title"] for s in info["slides"]]
    for t in powerpoint_tools.REQUIRED_TITLES:
        assert any(x.startswith(t) for x in titles), t
    assert all("SIMULATED" in s["text"] for s in info["slides"])
    assert "https://co1.test" in "\n".join(s["text"] for s in info["slides"] if s["title"].startswith("References"))


def test_many_competitors_paginate_tables_and_references(tmp_path):
    recs = [r.model_copy(update={"name": f"Co{i}", "website": f"https://co{i}.test"}) for i in range(10) for r in [records()[0]]]
    for i, r in enumerate(recs):
        r.sources = [SourceRef(url=f"https://co{i}.test/p{k}", title="t", retrieved_at=datetime.now(timezone.utc)) for k in range(2)]
    p = tmp_path / "big.pptx"
    powerpoint_tools.build_presentation(p, recs, {"industry": "X", "simulated": False})
    titles = [s["title"] for s in powerpoint_tools.inspect_presentation(p)["slides"]]
    assert sum(t.startswith("Competitor comparison") for t in titles) == 2
    assert sum(t.startswith("References") for t in titles) >= 2


def test_upload_validation_and_text_extraction(tmp_path):
    assert validate_upload("../../evil/report.PDF", 10, 5) == "report.PDF"  # directory part stripped
    for bad, size in [("x.exe", 10), ("x.pdf", 0), ("x.txt", 6 * 1024 * 1024), ("noext", 5)]:
        with pytest.raises(FileRejected):
            validate_upload(bad, size, 5)
    t = tmp_path / "a.txt"
    t.write_text("hello world")
    assert extract_text(t)["text"] == "hello world"
    import pymupdf as fitz

    doc = fitz.open()
    doc.new_page().insert_text((72, 72), "PDF content here")
    doc.save(tmp_path / "a.pdf")
    assert "PDF content here" in extract_text(tmp_path / "a.pdf")["text"]
    (tmp_path / "t.csv").write_text("a,b\n1,2\n")
    assert extract_text(tmp_path / "t.csv")["rows"] == 1
    (tmp_path / "broken.pdf").write_bytes(b"not a pdf")
    with pytest.raises(FileRejected):
        extract_text(tmp_path / "broken.pdf")


def test_safe_join_blocks_traversal(tmp_path):
    assert safe_join(tmp_path, "a", "b.txt").parent.name == "a"
    for bad in ("../x", "a/../../x", "/etc/passwd"):
        with pytest.raises(UnsafePath):
            safe_join(tmp_path, bad)


async def test_upload_endpoint_and_input_document_flows_into_run(client):
    bad = await client.post("/api/uploads", files={"file": ("evil.exe", b"MZ...", "application/octet-stream")})
    assert bad.status_code == 422
    up = await client.post("/api/uploads", files={"file": ("seed.txt", b"Known players: Foo, Bar. " * 10, "text/plain")})
    assert up.status_code == 201
    r = (await client.post("/api/objectives", json={"objective": OBJECTIVE, "mode": "autonomous", "upload_ids": [up.json()["id"]]})).json()
    assert r["tasks"][0]["action"] == "file.read_inputs"
    final = await wait_status(client, r["id"], {"completed", "failed", "partial", "halted"})
    assert final["status"] == "completed"
    t1 = (await client.get(f"/api/runs/{r['id']}/tasks/T1")).json()
    assert t1["output"]["files"][0]["name"] == "seed.txt"
    pptx = next(a for a in final["artifacts"] if a["kind"] == "pptx")
    from pptx import Presentation

    prs = Presentation(io.BytesIO((await client.get(pptx["download_url"])).content))
    assert "seed.txt" in " ".join(sh.text_frame.text for s in prs.slides for sh in s.shapes if sh.has_text_frame)


async def test_download_authorization(client, tmp_path):
    r = (await client.post("/api/objectives", json={"objective": OBJECTIVE, "mode": "autonomous"})).json()
    final = await wait_status(client, r["id"], {"completed"})
    art = next(a for a in final["artifacts"] if a["kind"] == "xlsx")
    assert (await client.get(art["download_url"])).status_code == 200
    for bad in ("art_000000000000", "..%2F..%2Fetc%2Fpasswd", "art_../x", "x" * 70):
        assert (await client.get(f"/api/artifacts/{bad}/download")).status_code in (404, 422)
    # a DB row pointing outside its root is refused
    with session_scope() as s:
        a = s.get(Artifact, art["id"])
        a.rel_path = "../../../../Windows/win.ini"
    assert (await client.get(art["download_url"])).status_code == 404
    with session_scope() as s:
        s.get(Artifact, art["id"]).rel_path = f"{r['id']}/competitors.xlsx"
    assert (await client.get(art["download_url"])).status_code == 200
    # a tampered file fails its integrity check and is not served
    from app.services.artifact_service import resolve_download

    path, _ = resolve_download(art["id"])
    path.write_bytes(path.read_bytes() + b"tamper")
    assert (await client.get(art["download_url"])).status_code == 404


async def test_verification_detects_tampered_deliverable(client):
    """Independent verification re-reads the file: a workbook edited to disagree with the evidence must fail."""
    from app.agents import verification_agent as va
    from app.agents.base import AgentContext
    from app.agents.orchestrator import Orchestrator, RunHandle
    from app.services import artifact_service, runtime_settings, store
    from app.services.llm import build_providers

    r = (await client.post("/api/objectives", json={"objective": OBJECTIVE, "mode": "autonomous"})).json()
    final = await wait_status(client, r["id"], {"completed"})
    art = next(a for a in final["artifacts"] if a["kind"] == "xlsx")
    path, _ = artifact_service.resolve_download(art["id"])
    wb = load_workbook(path)
    ws = wb["Comparison"]
    for row in ws.iter_rows(min_row=2):
        if row[6].value is not None:
            row[6].value = 1.0  # a price that no evidence supports
            break
    wb.save(path)
    orch = Orchestrator(r["id"], RunHandle(r["id"]))
    tasks = store.get_tasks(r["id"])
    outputs = orch.outputs()
    tv = next(t for t in tasks if t.action == "workspace.build_excel")
    ctx = AgentContext(run_id=r["id"], simulated=True, providers=build_providers(True, "x"), cancel=RunHandle(r["id"]).cancel,
                       settings=runtime_settings.load(), run_params={}, task=tv, outputs=outputs, tasks=tasks)
    recs = va._records_from(ctx)
    checks = va.excel_checks(path, recs, ctx)
    failed = {c.name for c in checks if not c.passed}
    assert "xlsx.prices_match_evidence" in failed
