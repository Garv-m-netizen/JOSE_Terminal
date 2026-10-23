"""Templates written as 'Student Name: [STUDENT NAME]' (the common school report-card style)."""
from __future__ import annotations

import io
import json

import docx
import pymupdf
import pytest

from app.tools import template_fill as tf
from app.tools.inspect_files import inspect, norm_key

from .test_custom_workflow import UPLOADS, finish, go, marks_xlsx, plan_json, planner_only


def bracket_docx(path):
    d = docx.Document()
    d.add_paragraph("[SCHOOL NAME]")
    p = d.add_paragraph()
    p.add_run("Total Marks: ").bold = True  # formatting must survive
    p.add_run("[____ / 400]")
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text = "Student Name: [STUDENT NAME]"
    t.cell(0, 1).text = "Father's Name: [FATHER'S NAME]"
    t.cell(1, 0).text = "Roll Number: [ROLL NO.]"
    t.cell(1, 1).text = "Date of Birth: [DD/MM/YYYY]"
    r = t.cell(0, 0).paragraphs[0]
    r.clear()
    r.add_run("Student Name: ").bold = True
    r.add_run("[STUDENT NAME]")  # label and bracket in different runs
    d.save(str(path))


def test_key_normalisation():
    assert norm_key("Father's Name:") == norm_key("[FATHER'S NAME]")
    assert norm_key("Student Name:") == norm_key("[STUDENT NAME]") == norm_key("{{student_name}}") == "studentname"
    assert norm_key("Roll No.") == norm_key("[ROLL NO.]")


def test_inspection_lists_bracket_fields(tmp_path):
    bracket_docx(tmp_path / "t.docx")
    i = inspect(tmp_path / "t.docx", "t.docx")
    labels = {f["label"] for f in i["bracket_fields"]}
    assert {"Student Name", "Father's Name", "Roll Number", "Date of Birth", "Total Marks"} <= labels
    assert "STUDENT NAME" in i["bracket_tokens"] and "SCHOOL NAME" in i["bracket_tokens"]


@pytest.mark.parametrize("keys", [
    {"Student Name": "Aarav Sharma", "Father's Name": "Rakesh Sharma"},            # label
    {"Student Name:": "Aarav Sharma", "Father's Name:": "Rakesh Sharma"},          # label with colon
    {"STUDENT NAME": "Aarav Sharma", "FATHER'S NAME": "Rakesh Sharma"},            # bracket text
    {"[Student Name]": "Aarav Sharma", "[Father's Name]": "Rakesh Sharma"},        # with brackets
])
def test_docx_bracket_fill_in_every_key_style(tmp_path, keys):
    bracket_docx(tmp_path / "t.docx")
    rep = tf.fill_docx(tmp_path / "t.docx", tmp_path / "o.docx", keys)
    assert rep["unfilled"] == []
    text = tf.read_docx_text(tmp_path / "o.docx")
    assert "Student Name: Aarav Sharma" in text and "Father's Name: Rakesh Sharma" in text
    assert "[STUDENT NAME]" not in text and "[FATHER'S NAME]" not in text
    assert "[ROLL NO.]" in text and "[SCHOOL NAME]" in text and "[____ / 400]" in text  # fields the user did not ask for stay untouched


def test_formatting_of_labels_is_preserved(tmp_path):
    bracket_docx(tmp_path / "t.docx")
    tf.fill_docx(tmp_path / "t.docx", tmp_path / "o.docx", {"Student Name": "Aarav Sharma", "Total Marks": "380 / 400"})
    d = docx.Document(str(tmp_path / "o.docx"))
    para = d.tables[0].cell(0, 0).paragraphs[0]
    assert [r.text for r in para.runs] == ["Student Name: ", "Aarav Sharma"] and para.runs[0].bold
    assert d.paragraphs[1].runs[0].bold and d.paragraphs[1].text == "Total Marks: 380 / 400"


def test_unknown_field_is_reported_not_silently_ignored(tmp_path):
    bracket_docx(tmp_path / "t.docx")
    rep = tf.fill_docx(tmp_path / "t.docx", tmp_path / "o.docx", {"Student Name": "A", "Mother's Name": "B"})
    assert rep["unfilled"] == ["Mother's Name"]


def test_pdf_bracket_fill_keeps_borders_and_other_fields(tmp_path):
    doc = pymupdf.open()
    pg = doc.new_page()
    pg.draw_rect(pymupdf.Rect(60, 100, 540, 150))
    pg.insert_text((70, 120), "Student Name: [STUDENT NAME]", fontsize=11)
    pg.insert_text((70, 140), "Roll Number: [ROLL NO.]", fontsize=11)
    pg.insert_text((300, 120), "Father's Name: [FATHER'S NAME]", fontsize=11)
    doc.save(tmp_path / "t.pdf")
    rep = tf.fill_pdf(tmp_path / "t.pdf", tmp_path / "o.pdf", {"Student Name": "Aarav Sharma", "Father's Name": "Rakesh Sharma"})
    assert rep["unfilled"] == []
    text = tf.read_pdf_text(tmp_path / "o.pdf")
    assert "Aarav Sharma" in text and "Rakesh Sharma" in text and "[STUDENT NAME]" not in text and "[ROLL NO.]" in text
    assert len(pymupdf.open(tmp_path / "o.pdf")[0].get_drawings()) >= 1  # the table border is still there


async def test_end_to_end_with_a_bracket_template(client, monkeypatch, tmp_path):
    bracket_docx(tmp_path / "report_card_template.docx")
    plan = plan_json(template="report_card_template.docx", mapping={"Student Name:": "Student Name", "Father's Name:": "Father's Name"})
    uploads = [("marks.xlsx", marks_xlsx()), ("report_card_template.docx", (tmp_path / "report_card_template.docx").read_bytes())]
    r, spy = await go(client, monkeypatch, planner_only(plan), uploads)
    assert any("bracket_fields" in p and "STUDENT NAME" in p for _, p in spy.calls)  # the planner was shown the bracket fields
    final = await finish(client, r["id"])
    assert final["status"] == "completed", final["error"]
    d = docx.Document(io.BytesIO((await client.get(final["artifacts"][0]["download_url"])).content))
    text = "\n".join(c.text for t in d.tables for row in t.rows for c in row.cells)
    assert "Aarav Sharma" in text and "Rakesh Sharma" in text and "[STUDENT NAME]" not in text


async def test_error_lists_what_the_template_really_contains(client, monkeypatch, tmp_path):
    bracket_docx(tmp_path / "report_card_template.docx")
    bad = plan_json(template="report_card_template.docx", mapping={"Pupil": "Student Name"})
    seen = {}

    def h(name, payload, prompt):
        if name == "plan_workflow":
            return bad
        seen["prompt"] = prompt
        return {"args_json": json.dumps({"template": "report_card_template.docx", "source": "@T2", "mapping": {"Student Name": "Student Name"},
                                         "filename_pattern": "RC_{Student Name}", "output_format": "same"})}

    uploads = [("marks.xlsx", marks_xlsx()), ("report_card_template.docx", (tmp_path / "report_card_template.docx").read_bytes())]
    r, _ = await go(client, monkeypatch, h, uploads)
    final = await finish(client, r["id"])
    assert final["status"] == "completed"
    assert "Fields this template actually contains" in seen["prompt"] and "STUDENT NAME" in seen["prompt"]
