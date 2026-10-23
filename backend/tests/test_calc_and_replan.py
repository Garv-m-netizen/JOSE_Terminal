from __future__ import annotations

import io
import json
from types import SimpleNamespace

import docx
import httpx
import openpyxl
import pytest

from app.agents import custom_planner
from app.agents.custom_planner import coverage_hints
from app.config import get_settings
from app.schemas.domain import TaskSpec
from app.services.llm import GeminiLLM, GroqLLM, LLMError, _groq_model_cache
from app.tools.inspect_files import inspect
from app.tools.safe_expr import ExprError, compile_expr, evaluate
from app.workflows.custom import validate_custom_plan
from app.workflows.research_report import PlanError

from .test_custom_workflow import finish, go, planner_only
from .test_bracket_templates import bracket_docx


def calc(expr, row):
    tree, cols = compile_expr(expr)
    return evaluate(tree, cols, lambda c: row[c])


def test_safe_expressions():
    r = {"Maths": 88, "Science": "91", "Name": "Diya", "P": 0.5}
    assert calc("sum([Maths],[Science])", r) == 179  # numeric strings are coerced
    assert calc("round(([Maths]+[Science])/200*100, 2)", r) == 89.5
    assert calc("str(round([Maths]/4, 1)) + ' %'", r) == "22 %"
    assert calc("'PASS' if [Maths] >= 33 else 'FAIL'", r) == "PASS"
    assert calc("[Name]", r) == "Diya" and calc("upper([Name]) + '-' + str([Maths])", r) == "DIYA-88"
    assert calc("avg([Maths],[Science])", r) == 89.5 and calc("max([Maths],[Science])", r) == 91


@pytest.mark.parametrize("bad", [
    "__import__('os').system('dir')", "[Maths].__class__", "open('x')", "eval('1')", "[Maths] if", "lambda: 1", "[x for x in [1]]",
    "exec('1')", "getattr([Maths],'a')", "(1).real", "'a'*10**9", "x + 1", "1" * 500,
])
def test_expression_sandbox_rejects_everything_else(bad):
    with pytest.raises((ExprError, ValueError)):
        tree, cols = compile_expr(bad)
        evaluate(tree, cols, lambda c: 1)


def test_expression_errors_are_clear():
    with pytest.raises(ExprError, match="division by zero"):
        calc("[A]/[B]", {"A": 1, "B": 0})
    with pytest.raises(ExprError, match="not a number"):
        calc("[A]*2", {"A": "abc"})
    with pytest.raises(ExprError, match="empty"):
        calc("[A]*2", {"A": None})


def workbook_bytes():
    wb = openpyxl.Workbook()
    a = wb.active
    a.title = "Student Records"
    a.append(["Student ID", "Student Name", "Father's Name"])
    for r in (("S1", "Aarav Sharma", "Rajesh Sharma"), ("S2", "Ananya Verma", "Sanjay Verma"), ("S3", "Vihaan Singh", "Manoj Singh")):
        a.append(r)
    m = wb.create_sheet("Subject Marks")
    m.append(["Student ID", "Student Name", "Maths", "Science", "SST", "English", "Total / 400", "Percentage"])
    for r in (("S1", "Aarav Sharma", 88, 91, 84, 90, 353, 0.8825), ("S2", "Ananya Verma", 95, 89, 92, 94, 370, 0.925), ("S3", "Vihaan Singh", 76, 82, 79, 85, 322, 0.805)):
        m.append(r)
    b = io.BytesIO()
    wb.save(b)
    return b.getvalue()


def card_with_grid(path):
    d = docx.Document()
    t0 = d.add_table(rows=2, cols=2)
    t0.cell(0, 0).text, t0.cell(0, 1).text = "Student Name: [STUDENT NAME]", "Father's Name: [FATHER'S NAME]"
    t0.cell(1, 0).text, t0.cell(1, 1).text = "Roll Number: [ROLL NO.]", "Date of Birth: [DD/MM/YYYY]"
    t = d.add_table(rows=5, cols=3)
    for j, h in enumerate(("Subject", "Maximum Marks", "Marks Obtained")):
        t.cell(0, j).text = h
    for i, s in enumerate(("Maths", "Science", "SST", "English"), 1):
        t.cell(i, 0).text, t.cell(i, 1).text, t.cell(i, 2).text = s, "[100]", "[____]"
    d.add_paragraph("Total Marks: [____ / 400]")
    d.add_paragraph("Percentage: [____ %]")
    d.save(str(path))


def full_plan(extra_map=None, compute_cols=None):
    mapping = {"STUDENT NAME": "Student Name", "FATHER'S NAME": "Father's Name", "ROLL NO.": "Roll Number", "Total Marks": "Total Marks", "Percentage": "Percentage"}
    for s in ("Maths", "Science", "SST", "English"):
        mapping[f"{s} / Marks Obtained"] = s
    mapping.update(extra_map or {})
    cols = compute_cols or {"Total Marks": "sum([Maths],[Science],[SST],[English])", "Percentage": "round([Total Marks]/400*100, 2)", "Roll Number": "[Student ID]"}
    T = lambda k, tool, args, **kw: {"key": k, "title": f"{tool} step", "tool": tool, "args_json": json.dumps(args), **kw}  # noqa: E731
    return {"feasible": True, "assumptions": ["Student ID is used as the roll number"], "tasks": [
        T("T1", "table.read", {"file": "records.xlsx", "sheet": "Student Records"}),
        T("T2", "table.read", {"file": "records.xlsx", "sheet": "Subject Marks"}),
        T("T3", "table.join", {"left": "@T1", "right": "@T2", "on": "Student ID"}),
        T("T4", "table.compute", {"source": "@T3", "columns": cols, "check_against": {"Total Marks": "Total / 400"}}),
        T("T5", "table.select", {"source": "@T4", "first": 2}),
        T("T6", "template.fill", {"template": "card.docx", "source": "@T5", "mapping": mapping, "filename_pattern": "Card_{Student Name}", "output_format": "same"}),
    ]}


@pytest.fixture
def uploads(tmp_path):
    card_with_grid(tmp_path / "card.docx")
    return [("records.xlsx", workbook_bytes()), ("card.docx", (tmp_path / "card.docx").read_bytes())]


async def test_two_sheets_joined_calculated_and_filled_into_grid_cells(client, monkeypatch, uploads):
    r, spy = await go(client, monkeypatch, planner_only(full_plan()), uploads, objective="fill report cards with marks, total and percentage")
    assert r["status"] != "rejected", r["error"]
    assert any("Subject Marks" in p and "sheet_details" in p and "Maths / Marks Obtained" in p for _, p in spy.calls)  # planner sees every sheet + grid cells
    final = await finish(client, r["id"])
    assert final["status"] == "completed", final["error"]
    assert sorted(a["filename"] for a in final["artifacts"]) == ["Card_Aarav Sharma.docx", "Card_Ananya Verma.docx"]
    d = docx.Document(io.BytesIO((await client.get(final["artifacts"][1]["download_url"])).content))
    cells = [c.text for t in d.tables for row in t.rows for c in row.cells]
    text = " | ".join(cells) + " | " + " | ".join(p.text for p in d.paragraphs)
    for expect in ("Ananya Verma", "Sanjay Verma", "S2", "95", "89", "92", "94", "Total Marks: 370", "Percentage: 92.5"):
        assert expect in text, expect
    assert "Vihaan" not in text and "[STUDENT NAME]" not in text
    v = (await client.get(f"/api/runs/{r['id']}/verification")).json()
    passed = {c["check"] for c in v["current"] if c["passed"]}
    assert {"join.row_count", "compute.recomputed_equal", "compute.vs_existing_columns", "compute.matches_sheet", "fill.values_present"} <= passed


async def test_wrong_calculation_is_caught_against_the_sheet(client, monkeypatch, uploads):
    bad = full_plan(compute_cols={"Total Marks": "sum([Maths],[Science],[SST])", "Percentage": "round([Total Marks]/400*100, 2)", "Roll Number": "[Student ID]"})
    r, _ = await go(client, monkeypatch, planner_only(bad), uploads)
    final = await finish(client, r["id"])
    assert final["status"] != "completed" and not final["artifacts"]
    v = (await client.get(f"/api/runs/{r['id']}/verification")).json()
    fail = next(c for c in v["checks"] if c["check"] == "compute.matches_sheet" and not c["passed"])
    assert "Total Marks" in fail["reason"] and "Total / 400" in fail["reason"]


def spec(key, tool, args, deps=(), agent="workspace"):
    return TaskSpec(key=key, title=f"{tool} step", agent=agent, action=f"tool.{tool}", params={"args": args}, depends_on=list(deps))


def test_plan_validation_rejects_text_output_used_as_a_table():
    plan = [spec("T1", "table.read", {"file": "m.xlsx"}), spec("T2", "llm.generate", {"instruction": "calculate totals", "input": "@T1"}, ["T1"], "research"),
            spec("T3", "template.fill", {"template": "c.docx", "source": "@T2", "mapping": {"a": "b"}}, ["T2"])]
    with pytest.raises(PlanError, match="must reference a whole table"):
        validate_custom_plan(plan, {"m.xlsx", "c.docx"})
    validate_custom_plan([plan[0], spec("T2", "table.compute", {"source": "@T1", "columns": {"x": "[A]*2"}}, ["T1"]),
                          spec("T3", "template.fill", {"template": "c.docx", "source": "@T2", "mapping": {"a": "x"}}, ["T2"])], {"m.xlsx", "c.docx"})


async def test_planner_is_corrected_when_it_uses_the_text_tool_for_maths(client, monkeypatch, uploads):
    wrong = full_plan()
    wrong["tasks"][3] = {"key": "T4", "title": "calculate", "tool": "llm.generate", "args_json": json.dumps({"instruction": "add the marks", "input": "@T3"})}
    state = {"n": 0}

    def h(name, payload, prompt):
        state["n"] += 1
        return wrong if state["n"] == 1 else full_plan()

    r, spy = await go(client, monkeypatch, h, uploads)
    assert r["status"] != "rejected" and [t["action"] for t in r["tasks"]][3] == "tool.table.compute"
    assert "must reference a whole table" in spy.calls[1][1]


async def test_failed_arguments_escalate_to_a_full_replan(client, monkeypatch, uploads):
    broken = full_plan(extra_map={"Nonexistent Field": "Student Name"})
    state = {"revise": 0, "plans": 0}

    def h(name, payload, prompt):
        if name == "plan_workflow":
            state["plans"] += 1
            return broken if state["plans"] == 1 else full_plan()
        state["revise"] += 1  # the argument-level repair keeps returning the same broken mapping
        return {"args_json": json.dumps(json.loads(broken["tasks"][5]["args_json"])), "explanation": "unchanged"}

    monkeypatch.setenv("MAX_RETRIES", "3")
    r, spy = await go(client, monkeypatch, h, uploads)
    final = await finish(client, r["id"])
    assert final["status"] == "completed", final["error"]
    assert state["revise"] == 1 and state["plans"] == 2
    assert any("FAILED while running" in p and "Nonexistent Field" in p for _, p in spy.calls)
    trace = (await client.get(f"/api/runs/{r['id']}/trace")).json()
    assert any(e["type"] == "plan.updated" and e["details"].get("reason") == "full replan after failure" for e in trace["events"])


def test_coverage_hint_finds_unmapped_cells_with_matching_columns(tmp_path):
    card_with_grid(tmp_path / "card.docx")
    files = [{"name": "records.xlsx", "kind": "table", "columns": ["Student ID"], "sheet_details": [{"sheet": "Subject Marks", "columns": ["Maths", "Science", "SST", "English"]}]},
             {**inspect(tmp_path / "card.docx", "card.docx")}]
    partial = custom_planner.to_specs(custom_planner.CustomPlan.model_validate(full_plan(extra_map={}) | {"tasks": full_plan()["tasks"]}))
    fill = next(t for t in partial if t.action == "tool.template.fill")
    for s in ("SST", "English"):
        fill.params["args"]["mapping"].pop(f"{s} / Marks Obtained")
    hints = coverage_hints(partial, files)
    assert any("SST / Marks Obtained" in h and "'SST'" in h for h in hints) and any("English" in h for h in hints) and not any("Maths" in h for h in hints)


# ------------------------------------------------------------------ provider resilience

async def test_gemini_retries_temporary_overload_before_giving_up(monkeypatch):
    monkeypatch.setenv("GEMINI_API_KEY", "x" * 12)
    get_settings.cache_clear()
    from app.schemas.domain import CandidateList

    calls = {"n": 0}

    class Models:
        async def generate_content(self, model, contents, config):
            calls["n"] += 1
            if calls["n"] < 3:
                raise RuntimeError("503 UNAVAILABLE. This model is currently experiencing high demand")
            return SimpleNamespace(function_calls=[SimpleNamespace(args={"competitors": [{"name": "A", "url": "https://a.test"}]})])

    llm = GeminiLLM.__new__(GeminiLLM)
    llm._client, llm._model = SimpleNamespace(aio=SimpleNamespace(models=Models())), "m"
    assert (await llm.call("f", "d", CandidateList, "s", "p")).competitors[0].name == "A" and calls["n"] == 3
    calls["n"] = -99  # a permanent error is not retried
    class Bad(Models):
        async def generate_content(self, *a, **k):
            calls["n"] += 1
            raise RuntimeError("400 INVALID_ARGUMENT")
    llm._client = SimpleNamespace(aio=SimpleNamespace(models=Bad()))
    calls["n"] = 0
    with pytest.raises(LLMError):
        await llm.call("f", "d", CandidateList, "s", "p")
    assert calls["n"] == 1


async def test_groq_picks_a_model_the_account_can_use_and_waits_out_rate_limits(monkeypatch):
    from app.schemas.domain import CandidateList

    monkeypatch.setenv("GROQ_API_KEY", "gsk_x" * 4)
    monkeypatch.setenv("GROQ_MODEL", "llama-3.3-70b-versatile")
    get_settings.cache_clear()
    _groq_model_cache.clear()
    seen, state = [], {"limited": False}

    def handler(req: httpx.Request):
        if req.method == "GET":
            return httpx.Response(200, json={"data": [{"id": "whisper-large-v3"}, {"id": "qwen/qwen3.8-27b"}, {"id": "openai/gpt-oss-120b"}]})
        body = json.loads(req.content)
        seen.append(body["model"])
        if body["model"] == "llama-3.3-70b-versatile":
            return httpx.Response(404, json={"error": {"code": "model_not_found", "message": "The model does not exist or you do not have access"}})
        if not state["limited"]:
            state["limited"] = True
            return httpx.Response(429, text='{"error":{"message":"Rate limit reached ... Please try again in 0.5s."}}')
        return httpx.Response(200, json={"choices": [{"message": {"tool_calls": [{"function": {"arguments": json.dumps({"competitors": [{"name": "A", "url": "https://a.test"}]})}}]}}]})

    monkeypatch.setenv("LLM_RETRY_SCALE", "0")
    get_settings.cache_clear()
    llm = GroqLLM(transport=httpx.MockTransport(handler))
    out = await llm.call("f", "d", CandidateList, "s", "p")
    assert out.competitors[0].name == "A"
    assert seen == ["llama-3.3-70b-versatile", "qwen/qwen3.8-27b", "qwen/qwen3.8-27b"]  # skipped the speech model, switched, then retried after the 429
    assert _groq_model_cache["llama-3.3-70b-versatile"] == "qwen/qwen3.8-27b"
    _groq_model_cache.clear()
