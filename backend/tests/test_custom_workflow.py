from __future__ import annotations

import asyncio
import io
import json

import docx
import openpyxl
import pymupdf
import pytest

from app.agents import custom_planner
from app.agents.intent_engine import parse_rules
from app.schemas.domain import TaskSpec
from app.services.llm import Providers, set_providers_override
from app.tools import template_fill as tf
from app.tools.inspect_files import inspect, read_table
from app.workflows.custom import RefError, collect_refs, resolve_args, validate_custom_plan
from app.workflows.research_report import PlanError

from .helpers import FakeLLM
from .test_exam_workflow import FakeTelegram

OBJECTIVE = ("see the attached excel file and take its name and fathers name and put those details in the attached report card template "
             "and make the report card of the first two students as pdfs")


# ------------------------------------------------------------------ fixtures

def marks_xlsx() -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Class 10-A Marks Sheet"])
    ws.append([])
    ws.append(["Roll No", "Student Name", "Father's Name", "Maths", "Science"])
    for r in ([1, "Aarav Sharma", "Rakesh Sharma", 91, 88.0], [2, "Diya Verma", "Anil Verma", 85, 90], [3, "Kabir Singh", "Harjit Singh", 70, 65]):
        ws.append(r)
    b = io.BytesIO()
    wb.save(b)
    return b.getvalue()


def card_docx() -> bytes:
    d = docx.Document()
    d.add_heading("REPORT CARD", 1)
    d.add_paragraph("Name of Student: {{Student Name}}")
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text, t.cell(1, 0).text, t.cell(1, 1).text = "Father's Name:", "Class:", "X-A"
    b = io.BytesIO()
    d.save(b)
    return b.getvalue()


def card_pdf() -> bytes:
    doc = pymupdf.open()
    pg = doc.new_page()
    pg.insert_text((72, 90), "REPORT CARD", fontsize=18)
    pg.insert_text((72, 130), "Student: {{Student Name}}", fontsize=12)
    pg.insert_text((72, 160), "Father's Name:", fontsize=12)
    return doc.tobytes()


async def upload(client, name, data):
    r = await client.post("/api/uploads", files={"file": (name, data, "application/octet-stream")})
    assert r.status_code == 201, r.text
    return r.json()["id"]


def plan_json(output_format="same", template="report_card.docx", mapping=None, extra_tasks=None, first=2):
    mapping = mapping or {"Student Name": "Student Name", "Father's Name": "Father's Name"}
    tasks = [
        {"key": "T1", "title": "Read the marks sheet", "tool": "table.read", "args_json": json.dumps({"file": "marks.xlsx"}), "expected_output": "rows"},
        {"key": "T2", "title": "Keep the first two students", "tool": "table.select", "args_json": json.dumps({"source": "@T1", "first": first})},
        {"key": "T3", "title": "Fill the report card template", "tool": "template.fill", "expected_output": "one file per student",
         "args_json": json.dumps({"template": template, "source": "@T2", "mapping": mapping, "filename_pattern": "ReportCard_{Student Name}", "output_format": output_format})},
        *(extra_tasks or []),
    ]
    return {"feasible": True, "assumptions": ["'first two' means the first two data rows"], "tasks": tasks}


class Spy:
    def __init__(self, handler):
        self.handler, self.calls = handler, []

    def __call__(self, name, payload, prompt):
        self.calls.append((name, prompt))
        return self.handler(name, payload, prompt)


async def go(client, monkeypatch, handler, uploads, tg=None, objective=OBJECTIVE, mode="autonomous"):
    monkeypatch.setenv("JOSE_MODE", "real")
    monkeypatch.setenv("GEMINI_API_KEY", "x" * 12)
    monkeypatch.setenv("TELEGRAM_CHAT_ID", "-1001234567890")
    from app.config import get_settings

    get_settings.cache_clear()
    spy = Spy(handler)
    prov = Providers(llm=FakeLLM(spy), search=None, fetcher=None, browser=None, simulated=False, messenger=tg)
    set_providers_override(lambda s, i: prov)
    ids = [await upload(client, n, d) for n, d in uploads]
    r = (await client.post("/api/objectives", json={"objective": objective, "mode": mode, "upload_ids": ids})).json()
    return r, spy


async def finish(client, run_id, decision="approved", timeout=40):
    end = asyncio.get_running_loop().time() + timeout
    while asyncio.get_running_loop().time() < end:
        run = (await client.get(f"/api/runs/{run_id}")).json()
        if run["status"] in ("completed", "failed", "partial", "halted", "cancelled", "rejected"):
            return run
        for a in run["pending_approvals"]:
            await client.post(f"/api/approvals/{a['id']}/resolve", json={"decision": decision})
        await asyncio.sleep(0.05)
    raise AssertionError("run did not finish")


UPLOADS = [("marks.xlsx", marks_xlsx()), ("report_card.docx", card_docx())]


# ------------------------------------------------------------------ units

def test_user_objective_is_routed_to_the_general_planner():
    i = parse_rules(OBJECTIVE, has_files=True)
    assert not i.feasible and i.workflow_type == "custom"
    assert parse_rules("Research five competitors in the AI automation market and create an Excel comparison").workflow_type == "research_report"
    assert parse_rules("make a maths test for class 10").workflow_type == "test_paper"
    assert parse_rules("Research five competitors in the CRM market and publish the findings").workflow_type == "research_report"  # still hard-rejected


def test_references():
    out = {"T1": {"rows": [{"a": 1}, {"a": 2}], "files": [{"artifact_id": "x1"}, {"artifact_id": "x2"}]}}
    assert resolve_args("@T1.files.*.artifact_id", out) == ["x1", "x2"]
    assert resolve_args({"k": ["@T1.rows.1.a", "plain"]}, out) == {"k": [2, "plain"]}
    assert collect_refs({"a": "@T3", "b": ["@T1.rows"]}) == ["T3", "T1"]
    with pytest.raises(RefError):
        resolve_args("@T1.nope", out)
    with pytest.raises(RefError):
        resolve_args("@T9", out)


def spec(key, tool, args, deps=(), agent="workspace"):
    return TaskSpec(key=key, title=f"{tool} step", agent=agent, action=f"tool.{tool}", params={"args": args}, depends_on=list(deps))


def test_plan_validation_rejects_bad_plans():
    ok = [spec("T1", "table.read", {"file": "marks.xlsx"}), spec("T2", "table.select", {"source": "@T1", "first": 2}, ["T1"])]
    validate_custom_plan(ok, {"marks.xlsx"})
    for bad, frag in [
        ([spec("T1", "shell.exec", {})], "allow-listed"),
        ([spec("T1", "table.read", {"file": "marks.xlsx", "extra": 1})], "invalid arguments"),
        ([ok[0], spec("T2", "table.select", {"source": "@T1"}, [])], "does not depend"),
        ([spec("T1", "table.read", {"file": "other.xlsx"})], "not attached"),
        ([spec("T1", "table.select", {"source": "@T2"}, ["T2"]), spec("T2", "table.select", {"source": "@T1"}, ["T1"])], "circular"),
        ([spec("T1", "table.read", {"file": "marks.xlsx"}, agent="browser")], "must be run by"),
    ]:
        with pytest.raises(PlanError, match=frag):
            validate_custom_plan(bad, {"marks.xlsx"})


def test_table_reader_finds_header_below_a_title(tmp_path):
    p = tmp_path / "m.xlsx"
    p.write_bytes(marks_xlsx())
    t = read_table(p)
    assert t["columns"] == ["Roll No", "Student Name", "Father's Name", "Maths", "Science"] and t["row_count"] == 3
    assert t["rows"][0]["Science"] == 88 and isinstance(t["rows"][0]["Science"], int)  # 88.0 shown as 88


def test_template_inspection_and_fill(tmp_path):
    (tmp_path / "c.docx").write_bytes(card_docx())
    (tmp_path / "c.pdf").write_bytes(card_pdf())
    d = inspect(tmp_path / "c.docx", "c.docx")
    assert d["placeholders"] == ["Student Name"] and "Father's Name" in d["labels"]
    rep = tf.fill_docx(tmp_path / "c.docx", tmp_path / "o.docx", {"student name": "Aarav", "Father's Name": "Rakesh"})  # case-insensitive keys
    assert rep["unfilled"] == [] and "Aarav" in tf.read_docx_text(tmp_path / "o.docx") and "{{" not in tf.read_docx_text(tmp_path / "o.docx")
    rep = tf.fill_pdf(tmp_path / "c.pdf", tmp_path / "o.pdf", {"Student Name": "Diya", "Father's Name": "Anil", "Missing": "x"})
    assert rep["unfilled"] == ["Missing"] and "Diya" in tf.read_pdf_text(tmp_path / "o.pdf") and "Anil" in tf.read_pdf_text(tmp_path / "o.pdf")
    wb = openpyxl.Workbook()
    wb.active.append(["Name:", None])
    wb.active.append(["{{Roll}}", "x"])
    wb.save(tmp_path / "t.xlsx")
    rep = tf.fill_xlsx(tmp_path / "t.xlsx", tmp_path / "o.xlsx", {"Name": "Kabir", "Roll": "3"})
    assert rep["unfilled"] == [] and "Kabir" in tf.read_xlsx_text(tmp_path / "o.xlsx")


# ------------------------------------------------------------------ end to end

def planner_only(plan):
    def h(name, payload, prompt):
        if name == "plan_workflow":
            return plan
        raise AssertionError(name)
    return h


async def test_report_cards_from_excel_and_word_template(client, monkeypatch):
    r, spy = await go(client, monkeypatch, planner_only(plan_json()), UPLOADS)
    assert r["status"] != "rejected", r["error"]
    assert r["workflow_type"] == "custom" and [t["action"] for t in r["tasks"]] == ["tool.table.read", "tool.table.select", "tool.template.fill"]
    assert "inventory" not in r and any("marks.xlsx" in p and "Father's Name" in p and "Student Name" in p for _, p in spy.calls)  # planner saw the real columns/placeholders
    final = await finish(client, r["id"])
    assert final["status"] == "completed", final["error"]
    files = sorted(a["filename"] for a in final["artifacts"])
    assert files == ["ReportCard_Aarav Sharma.docx", "ReportCard_Diya Verma.docx"]  # exactly the first two students
    for a in final["artifacts"]:
        d = docx.Document(io.BytesIO((await client.get(a["download_url"])).content))
        text = "\n".join(p.text for p in d.paragraphs) + "\n".join(c.text for t in d.tables for row in t.rows for c in row.cells)
        who = "Aarav Sharma" if "Aarav" in a["filename"] else "Diya Verma"
        father = "Rakesh Sharma" if "Aarav" in a["filename"] else "Anil Verma"
        assert who in text and father in text and "Kabir" not in text and "{{" not in text
    v = (await client.get(f"/api/runs/{r['id']}/verification")).json()
    assert {"fill.values_present", "fill.no_leftover_placeholders", "table.rows_match_file", "select.first_n"} <= {c["check"] for c in v["current"] if c["passed"]}
    assert final["intent"]["assumptions"] == ["'first two' means the first two data rows"]


async def test_report_cards_from_pdf_template(client, monkeypatch):
    plan = plan_json(template="card.pdf", output_format="same")
    r, _ = await go(client, monkeypatch, planner_only(plan), [("marks.xlsx", marks_xlsx()), ("card.pdf", card_pdf())])
    final = await finish(client, r["id"])
    assert final["status"] == "completed", final["error"]
    assert sorted(a["kind"] for a in final["artifacts"]) == ["pdf", "pdf"]
    text = pymupdf.open(stream=(await client.get(final["artifacts"][0]["download_url"])).content, filetype="pdf")[0].get_text()
    assert "Aarav Sharma" in text and "Rakesh Sharma" in text


@pytest.mark.skipif(tf.find_converter() is None, reason="no Word/LibreOffice installed")
async def test_word_template_converted_to_real_pdf(client, monkeypatch):
    r, _ = await go(client, monkeypatch, planner_only(plan_json(output_format="pdf")), UPLOADS)
    final = await finish(client, r["id"], timeout=120)
    assert final["status"] == "completed", final["error"]
    assert [a["kind"] for a in final["artifacts"]] == ["pdf", "pdf"]
    text = pymupdf.open(stream=(await client.get(final["artifacts"][1]["download_url"])).content, filetype="pdf")[0].get_text()
    assert "Diya Verma" in text and "Anil Verma" in text


async def test_wrong_column_name_is_repaired_by_the_planner(client, monkeypatch):
    bad = plan_json(mapping={"Student Name": "Pupil Name", "Father's Name": "Father's Name"})
    seen = {}

    def h(name, payload, prompt):
        if name == "plan_workflow":
            return bad
        if name == "revise_task_args":
            seen["prompt"] = prompt
            return {"args_json": json.dumps({"template": "report_card.docx", "source": "@T2", "mapping": {"Student Name": "Student Name", "Father's Name": "Father's Name"},
                                             "filename_pattern": "ReportCard_{Student Name}", "output_format": "same"}), "explanation": "column is 'Student Name'"}
        raise AssertionError(name)

    r, _ = await go(client, monkeypatch, h, UPLOADS)
    final = await finish(client, r["id"])
    assert final["status"] == "completed" and final["recovery_count"] == 1
    assert "Pupil Name" in seen["prompt"] and "available columns" in seen["prompt"]
    trace = (await client.get(f"/api/runs/{r['id']}/trace")).json()
    assert "recovery.started" in [e["type"] for e in trace["events"]]


async def test_unfixable_template_mapping_fails_without_producing_files(client, monkeypatch):
    bad = plan_json(mapping={"Grade Point": "Maths"})  # the template has no such field

    def h(name, payload, prompt):
        if name == "plan_workflow":
            return bad
        return {"args_json": json.dumps(bad["tasks"][2]["args_json"] and json.loads(bad["tasks"][2]["args_json"])), "explanation": "same"}

    r, _ = await go(client, monkeypatch, h, UPLOADS)
    final = await finish(client, r["id"])
    assert final["status"] in ("failed", "halted") and not final["artifacts"]
    assert "Grade Point" in final["error"] and final["recovery_count"] <= 3


async def test_infeasible_objective_is_refused_honestly(client, monkeypatch):
    def h(name, payload, prompt):
        return {"feasible": False, "reason": "There is no tool for sending email; I can only create files or post to Telegram.", "tasks": []}

    r, _ = await go(client, monkeypatch, h, UPLOADS, objective="take marks.xlsx and email each parent their child's report card")
    assert r["status"] == "rejected" and "no tool for sending email" in r["error"] and r["tasks"] == []


async def test_invalid_plan_from_the_model_is_sent_back_for_correction(client, monkeypatch):
    good, state = plan_json(), {"n": 0}
    bogus = {"feasible": True, "tasks": [{"key": "T1", "title": "run a script", "tool": "shell.exec", "args_json": "{}"}]}

    def h(name, payload, prompt):
        state["n"] += 1
        return bogus if state["n"] == 1 else good

    r, spy = await go(client, monkeypatch, h, UPLOADS)
    assert r["status"] != "rejected" and len(r["tasks"]) == 3
    assert "previous plan was invalid" in spy.calls[1][1] and "unknown tool 'shell.exec'" in spy.calls[1][1]


async def test_general_mode_needs_an_ai_model(client):
    ids = [await upload(client, n, d) for n, d in UPLOADS]
    r = (await client.post("/api/objectives", json={"objective": OBJECTIVE, "mode": "autonomous", "upload_ids": ids})).json()
    assert r["status"] == "rejected" and ("GEMINI_API_KEY" in r["error"] or "GROQ_API_KEY" in r["error"])


async def test_posting_generated_files_needs_approval_and_sends_only_run_files(client, monkeypatch):
    send = {"key": "T4", "title": "Post the report cards", "tool": "telegram.send_files",
            "args_json": json.dumps({"files": "@T3.files.*.artifact_id", "caption": "Report cards"})}
    tg = FakeTelegram()
    r, _ = await go(client, monkeypatch, planner_only(plan_json(extra_tasks=[send])), UPLOADS, tg=tg)
    assert r["tasks"][-1]["risk_level"] == "HIGH"
    run = {}
    for _ in range(300):
        run = (await client.get(f"/api/runs/{r['id']}")).json()
        if run["pending_approvals"]:
            break
        await asyncio.sleep(0.05)
    ap = run["pending_approvals"][0]
    assert tg.sent == [] and ap["risk_level"] == "HIGH" and "ReportCard_Aarav Sharma.docx" in ap["details"]["file"] and "Report cards" in ap["details"]["caption"]
    final = await finish(client, r["id"])
    assert final["status"] == "completed" and sorted(n for n, _ in tg.sent) == ["ReportCard_Aarav Sharma.docx", "ReportCard_Diya Verma.docx"]


async def test_denied_post_of_generated_files_sends_nothing(client, monkeypatch):
    send = {"key": "T4", "title": "Post", "tool": "telegram.send_files", "args_json": json.dumps({"files": "@T3.files.*.artifact_id"})}
    tg = FakeTelegram()
    r, _ = await go(client, monkeypatch, planner_only(plan_json(extra_tasks=[send])), UPLOADS, tg=tg)
    final = await finish(client, r["id"], "denied")
    assert final["status"] == "halted" and tg.sent == [] and len(final["artifacts"]) == 2


async def test_general_workflow_can_be_saved_and_rerun_with_new_files(client, monkeypatch):
    r, _ = await go(client, monkeypatch, planner_only(plan_json()), UPLOADS)
    assert (await finish(client, r["id"]))["status"] == "completed"
    wf = (await client.post("/api/workflows", json={"run_id": r["id"], "name": "Report cards"})).json()
    assert wf["workflow_type"] == "custom" and wf["params"]["files"] == ["marks.xlsx", "report_card.docx"] and len(wf["task_config"]) == 3
    assert (await client.post(f"/api/workflows/{wf['id']}/run", json={"params": {}, "mode": "autonomous"})).status_code == 422  # needs files
    ids = [await upload(client, n, d) for n, d in UPLOADS]
    r2 = (await client.post(f"/api/workflows/{wf['id']}/run", json={"params": {"upload_ids": ids}, "mode": "autonomous"})).json()
    assert (await finish(client, r2["id"]))["status"] == "completed"
    wrong = await upload(client, "other.xlsx", marks_xlsx())
    assert (await client.post(f"/api/workflows/{wf['id']}/run", json={"params": {"upload_ids": [wrong]}})).status_code == 422


async def test_plan_edit_in_guided_mode_is_validated(client, monkeypatch):
    r, _ = await go(client, monkeypatch, planner_only(plan_json()), UPLOADS, mode="guided")
    assert r["status"] == "awaiting_plan_approval"
    tasks = [dict(t) for t in r["tasks"]]
    tasks[2]["params"] = {"args": {"template": "missing.docx", "source": "@T2", "mapping": {"a": "b"}}}
    bad = await client.put(f"/api/runs/{r['id']}/plan", json={"tasks": tasks})
    assert bad.status_code == 422
    assert (await client.post(f"/api/runs/{r['id']}/plan/approve")).status_code == 200
    assert (await finish(client, r["id"]))["status"] == "completed"
