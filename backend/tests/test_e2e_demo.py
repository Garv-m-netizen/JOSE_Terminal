from __future__ import annotations

import io

from openpyxl import load_workbook
from pptx import Presentation

from .conftest import OBJECTIVE, wait_status


async def test_demo_autonomous_end_to_end(client):
    r = await client.post("/api/objectives", json={"objective": OBJECTIVE, "mode": "autonomous"})
    assert r.status_code == 201, r.text
    run = r.json()
    assert run["simulated"] is True
    final = await wait_status(client, run["id"], {"completed", "failed", "partial", "halted"})
    assert final["status"] == "completed", final.get("error")
    arts = (await client.get(f"/api/runs/{run['id']}/artifacts")).json()
    kinds = {a["kind"] for a in arts}
    assert {"xlsx", "pptx"} <= kinds
    for a in arts:
        d = await client.get(a["download_url"])
        assert d.status_code == 200 and len(d.content) > 1000
        if a["kind"] == "xlsx":
            wb = load_workbook(io.BytesIO(d.content))
            assert wb.sheetnames == ["Summary", "Comparison", "Sources"]
            assert wb["Comparison"].max_row == 6
        if a["kind"] == "pptx":
            assert len(Presentation(io.BytesIO(d.content)).slides) >= 6
    ver = (await client.get(f"/api/runs/{run['id']}/verification")).json()
    assert ver["all_passed"] and ver["passed"] > 20
